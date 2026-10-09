import math
import base64
import difflib
import io
import json
import re
import zipfile
from pathlib import Path

import pymupdf
from openpyxl import Workbook
from PIL import Image, ImageChops, ImageDraw, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Emu, Pt
from pypdf import PdfReader, PdfWriter
from pypdf.generic import RectangleObject


# ---------------------------------------------------------------
# Basic page tools
# ---------------------------------------------------------------
def parse_pages(spec: str, total: int) -> list[int]:
    """'1-3,5' -> [0, 1, 2, 4]. Raises ValueError on bad input."""
    pages = []
    for part in spec.replace(" ", "").split(","):
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-")
            start, end = int(a), int(b)
        else:
            start = end = int(part)
        if start < 1 or end < start or end > total:
            raise ValueError(f"Pages must be between 1 and {total}")
        pages.extend(range(start - 1, end))
    if not pages:
        raise ValueError("No pages selected")
    return pages


def merge_pdfs(streams: list, out: Path) -> None:
    writer = PdfWriter()
    for s in streams:
        writer.append(PdfReader(s))
    with out.open("wb") as fh:
        writer.write(fh)


def extract_pages(stream, spec: str, out: Path) -> None:
    reader = PdfReader(stream)
    writer = PdfWriter()
    for i in parse_pages(spec, len(reader.pages)):
        writer.add_page(reader.pages[i])
    with out.open("wb") as fh:
        writer.write(fh)


def rotate_pdf(stream, angle: int, spec: str, out: Path) -> None:
    reader = PdfReader(stream)
    writer = PdfWriter()
    total = len(reader.pages)
    targets = set(parse_pages(spec, total)) if spec.strip() else set(range(total))
    for i, page in enumerate(reader.pages):
        if i in targets:
            page.rotate(angle)  # clockwise
        writer.add_page(page)
    with out.open("wb") as fh:
        writer.write(fh)


def protect_pdf(stream, password: str, out: Path) -> None:
    reader = PdfReader(stream)
    if reader.is_encrypted:
        raise ValueError("This PDF is already password protected")
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt(user_password=password, algorithm="AES-256")
    with out.open("wb") as fh:
        writer.write(fh)


def unlock_pdf(stream, password: str, out: Path) -> None:
    reader = PdfReader(stream)
    if not reader.is_encrypted:
        raise ValueError("This PDF is not password protected")
    if reader.decrypt(password) == 0:
        raise ValueError("Wrong password")
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    with out.open("wb") as fh:
        writer.write(fh)


# ---------------------------------------------------------------
# Images -> PDF (JPG, PNG, WebP)
# ---------------------------------------------------------------
MAX_IMAGES = 50
PAGE_SIZES = {"a4": (595.0, 842.0), "letter": (612.0, 792.0)}
MARGINS = {"none": 0, "small": 20, "big": 50}


def _load_image(stream):
    try:
        img = Image.open(stream)
        if img.width * img.height > 50_000_000:
            raise ValueError("An image is too large")
        fmt = img.format
        img.load()
    except ValueError:
        raise
    except Exception:
        raise ValueError("One of the files is not a valid image")
    img = ImageOps.exif_transpose(img)  # fix sideways phone photos
    if img.mode in ("RGBA", "LA", "P"):
        img = img.convert("RGBA")
        bg = Image.new("RGB", img.size, "white")
        bg.paste(img, mask=img.getchannel("A"))
        img = bg
    elif img.mode != "RGB":
        img = img.convert("RGB")
    buf = io.BytesIO()
    if fmt == "PNG":
        img.save(buf, "PNG")
    else:
        img.save(buf, "JPEG", quality=95)
    return img.width, img.height, buf.getvalue()


def _add_image_page(doc, w, h, data, size, orientation, margin):
    if size == "fit":
        iw, ih = w * 0.75, h * 0.75  # pixels -> points at 96 dpi
        k = min(1, 14000 / max(iw, ih))  # PDF page size limit
        iw, ih = iw * k, ih * k
        pw, ph = iw + 2 * margin, ih + 2 * margin
        x, y = margin, margin
    else:
        pw, ph = PAGE_SIZES[size]
        landscape = (w > h) if orientation == "auto" else orientation == "landscape"
        if landscape:
            pw, ph = ph, pw
        s = min((pw - 2 * margin) / w, (ph - 2 * margin) / h)
        iw, ih = w * s, h * s
        x, y = (pw - iw) / 2, (ph - ih) / 2
    page = doc.new_page(width=pw, height=ph)
    page.insert_image(pymupdf.Rect(x, y, x + iw, y + ih), stream=data)


def images_to_pdf(streams: list, out: Path, size="a4", orientation="auto",
                  margin="small", separate=False) -> None:
    if not 1 <= len(streams) <= MAX_IMAGES:
        raise ValueError(f"Use between 1 and {MAX_IMAGES} images")
    if size not in (*PAGE_SIZES, "fit"):
        raise ValueError("Unknown page size")
    if orientation not in ("auto", "portrait", "landscape"):
        raise ValueError("Unknown orientation")
    if margin not in MARGINS:
        raise ValueError("Unknown margin")
    m = MARGINS[margin]
    items = [_load_image(s) for s in streams]

    if separate:
        with zipfile.ZipFile(out, "w", zipfile.ZIP_STORED) as z:
            for k, (w, h, data) in enumerate(items, start=1):
                doc = pymupdf.open()
                _add_image_page(doc, w, h, data, size, orientation, m)
                z.writestr(f"image_{k}.pdf", doc.tobytes(deflate=True))
                doc.close()
        return

    doc = pymupdf.open()
    try:
        for w, h, data in items:
            _add_image_page(doc, w, h, data, size, orientation, m)
        doc.save(out, garbage=3, deflate=True)
    finally:
        doc.close()


# ---------------------------------------------------------------
# PDF -> JPG / PNG
# ---------------------------------------------------------------
MAX_IMAGE_PAGES = 50


def pdf_to_images(src: Path, dpi: int, job_dir: Path, fmt: str = "jpg") -> Path:
    if fmt not in ("jpg", "png"):
        raise ValueError("Unknown image format")
    doc = pymupdf.open(src)
    try:
        if doc.needs_pass:
            raise ValueError("This PDF is password protected. Unlock it first")
        total = doc.page_count
        if total > MAX_IMAGE_PAGES:
            raise ValueError(f"Too many pages (max {MAX_IMAGE_PAGES}). Extract fewer pages first")
        matrix = pymupdf.Matrix(dpi / 72, dpi / 72)

        def render(page):
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            return pix.tobytes("png") if fmt == "png" else pix.tobytes("jpg", jpg_quality=90)

        if total == 1:
            out = job_dir / f"page_1.{fmt}"
            out.write_bytes(render(doc[0]))
            return out

        out = job_dir / "pages.zip"
        with zipfile.ZipFile(out, "w", zipfile.ZIP_STORED) as z:
            for i, page in enumerate(doc, start=1):
                z.writestr(f"page_{i}.{fmt}", render(page))
        return out
    finally:
        doc.close()


# ---------------------------------------------------------------
# PDF -> Excel
# ---------------------------------------------------------------
MAX_TABLE_PAGES = 100
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def pdf_to_excel(src: Path, out: Path) -> None:
    doc = pymupdf.open(src)
    try:
        if doc.needs_pass:
            raise ValueError("This PDF is password protected. Unlock it first")
        if doc.page_count > MAX_TABLE_PAGES:
            raise ValueError(f"Too many pages (max {MAX_TABLE_PAGES})")
        wb = Workbook()
        wb.remove(wb.active)
        found = 0
        for pno, page in enumerate(doc, start=1):
            for tno, table in enumerate(page.find_tables().tables, start=1):
                rows = table.extract()
                if not rows:
                    continue
                found += 1
                ws = wb.create_sheet(f"P{pno}_T{tno}")
                for r, row in enumerate(rows, start=1):
                    for c, value in enumerate(row, start=1):
                        if value is None:
                            continue
                        value = _ILLEGAL.sub("", str(value))
                        cell = ws.cell(row=r, column=c, value=value)
                        cell.data_type = "s"  # keep text as text, never as a formula
        if found == 0:
            raise ValueError(
                "No tables found. This tool needs a PDF with real tables "
                "(scanned PDFs are not supported)"
            )
        wb.save(out)
    finally:
        doc.close()


# ---------------------------------------------------------------
# PDF -> PowerPoint
# ---------------------------------------------------------------
MAX_SHAPES = 400  # keeps heavy vector pages from creating thousands of shapes


def _rgb_int(value: int) -> RGBColor:
    return RGBColor((value >> 16) & 255, (value >> 8) & 255, value & 255)


def _rgb_float(color) -> RGBColor:
    return RGBColor(*(int(round(c * 255)) for c in color[:3]))


def _in_boxes(bbox, boxes) -> bool:
    cx, cy = (bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2
    return any(b[0] <= cx <= b[2] and b[1] <= cy <= b[3] for b in boxes)


def _build_editable_slide(page, slide, left, top, scale):
    pt = 12700

    def size(v):
        return Emu(int(v * pt * scale))

    def X(x):
        return Emu(int(left + x * pt * scale))

    def Y(y):
        return Emu(int(top + y * pt * scale))

    tables = page.find_tables().tables
    table_boxes = [t.bbox for t in tables]
    blocks = page.get_text("dict")["blocks"]

    # 1. Images (bottom layer)
    for b in blocks:
        if b["type"] == 1 and b.get("image"):
            x0, y0, x1, y1 = b["bbox"]
            try:
                slide.shapes.add_picture(
                    io.BytesIO(b["image"]), X(x0), Y(y0), size(x1 - x0), size(y1 - y0)
                )
            except Exception:
                pass  # skip image formats PowerPoint can't take

    # 2. Shapes: filled rectangles and straight lines
    count = 0
    for d in page.get_drawings():
        if count >= MAX_SHAPES:
            break
        r = d["rect"]
        if _in_boxes(tuple(r), table_boxes):
            continue  # table borders are rebuilt as a real table below
        items = d["items"]
        fill, stroke = d.get("fill"), d.get("color")
        width = Pt(max((d.get("width") or 1) * scale, 0.5))

        if items and all(i[0] == "re" for i in items):
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, X(r.x0), Y(r.y0),
                size(max(r.width, 0.5)), size(max(r.height, 0.5)),
            )
            if fill:
                shape.fill.solid()
                shape.fill.fore_color.rgb = _rgb_float(fill)
            else:
                shape.fill.background()
            if stroke:
                shape.line.color.rgb = _rgb_float(stroke)
                shape.line.width = width
            else:
                shape.line.fill.background()
            shape.shadow.inherit = False
            count += 1
        elif len(items) == 1 and items[0][0] == "l" and stroke:
            p1, p2 = items[0][1], items[0][2]
            line = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT, X(p1.x), Y(p1.y), X(p2.x), Y(p2.y)
            )
            line.line.color.rgb = _rgb_float(stroke)
            line.line.width = width
            count += 1
        # curves, gradients and complex paths are skipped in this version

    # 3. Real tables
    for t in tables:
        rows = t.extract()
        if not rows:
            continue
        x0, y0, x1, y1 = t.bbox
        n_rows, n_cols = len(rows), max(len(r) for r in rows)
        graphic = slide.shapes.add_table(
            n_rows, n_cols, X(x0), Y(y0), size(x1 - x0), size(y1 - y0)
        )
        for r_i, row in enumerate(rows):
            for c_i in range(n_cols):
                value = row[c_i] if c_i < len(row) and row[c_i] is not None else ""
                cell = graphic.table.cell(r_i, c_i)
                cell.text = _ILLEGAL.sub("", str(value))
                for para in cell.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(max(10 * scale, 6))

    # 4. Text boxes (top layer), skipping text that is inside a table
    for b in blocks:
        if b["type"] != 0 or not b["lines"] or _in_boxes(b["bbox"], table_boxes):
            continue
        x0, y0, x1, y1 = b["bbox"]
        box = slide.shapes.add_textbox(
            X(x0), Y(y0), Emu(int(size(x1 - x0) * 1.05)), size(y1 - y0)
        )
        frame = box.text_frame
        frame.word_wrap = False
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        first = True
        for line in b["lines"]:
            para = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            for span in line["spans"]:
                text = _ILLEGAL.sub("", span["text"])
                if not text:
                    continue
                run = para.add_run()
                run.text = text
                run.font.size = Pt(max(span["size"] * scale, 4))
                run.font.bold = bool(span["flags"] & 16)
                run.font.italic = bool(span["flags"] & 2)
                run.font.name = span["font"].split("+")[-1]
                run.font.color.rgb = _rgb_int(span["color"])


def pdf_to_pptx(src: Path, mode: str, out: Path) -> None:
    doc = pymupdf.open(src)
    try:
        if doc.needs_pass:
            raise ValueError("This PDF is password protected. Unlock it first")
        if doc.page_count > MAX_IMAGE_PAGES:
            raise ValueError(f"Too many pages (max {MAX_IMAGE_PAGES})")

        pt = 12700  # EMU per point
        first = doc[0].rect
        slide_w, slide_h = first.width * pt, first.height * pt
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(int(slide_w)), Emu(int(slide_h))
        blank = prs.slide_layouts[6]

        for page in doc:
            rect = page.rect
            scale = min(slide_w / (rect.width * pt), slide_h / (rect.height * pt))
            left = (slide_w - rect.width * pt * scale) / 2
            top = (slide_h - rect.height * pt * scale) / 2
            slide = prs.slides.add_slide(blank)

            if mode == "image":
                zoom = 150 / 72
                pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
                slide.shapes.add_picture(
                    io.BytesIO(pix.tobytes("jpg", jpg_quality=85)),
                    Emu(int(left)), Emu(int(top)),
                    Emu(int(rect.width * pt * scale)), Emu(int(rect.height * pt * scale)),
                )
            else:
                _build_editable_slide(page, slide, left, top, scale)
        prs.save(out)
    finally:
        doc.close()


# ---------------------------------------------------------------
# Split and remove pages
# ---------------------------------------------------------------
MAX_SPLIT_FILES = 200


def _safe_pages(spec: str, total: int) -> list[int]:
    try:
        return parse_pages(spec, total)
    except ValueError as e:
        if "invalid literal" in str(e):
            raise ValueError("Use page numbers like 1-3,5")
        raise


def remove_pages(stream, spec: str, out: Path) -> None:
    reader = PdfReader(stream)
    if reader.is_encrypted:
        raise ValueError("This PDF is password protected. Unlock it first")
    total = len(reader.pages)
    targets = set(_safe_pages(spec, total))
    if len(targets) >= total:
        raise ValueError("You can't remove every page")
    writer = PdfWriter()
    for i, page in enumerate(reader.pages):
        if i not in targets:
            writer.add_page(page)
    with out.open("wb") as fh:
        writer.write(fh)


def split_pdf(stream, mode: str, value: str, out: Path) -> None:
    reader = PdfReader(stream)
    if reader.is_encrypted:
        raise ValueError("This PDF is password protected. Unlock it first")
    total = len(reader.pages)

    if mode == "every":
        groups = [[i] for i in range(total)]
    elif mode == "n":
        try:
            n = int(value)
        except ValueError:
            raise ValueError("Enter a whole number, for example 2")
        if n < 1:
            raise ValueError("The number must be 1 or more")
        groups = [list(range(s, min(s + n, total))) for s in range(0, total, n)]
    elif mode == "ranges":
        if not value.strip():
            raise ValueError("Enter ranges like 1-3,4-6,7")
        groups = [
            _safe_pages(part, total)
            for part in value.replace(" ", "").split(",")
            if part
        ]
    else:
        raise ValueError("Unknown split mode")

    if not groups:
        raise ValueError("Nothing to split")
    if len(groups) > MAX_SPLIT_FILES:
        raise ValueError(f"Too many files (max {MAX_SPLIT_FILES})")

    with zipfile.ZipFile(out, "w", zipfile.ZIP_STORED) as z:
        for k, pages in enumerate(groups, start=1):
            writer = PdfWriter()
            for i in pages:
                writer.add_page(reader.pages[i])
            buf = io.BytesIO()
            writer.write(buf)
            label = f"{pages[0] + 1}" if len(pages) == 1 else f"{pages[0] + 1}-{pages[-1] + 1}"
            z.writestr(f"part_{k}_pages_{label}.pdf", buf.getvalue())


# ---------------------------------------------------------------
# Organize PDF
# ---------------------------------------------------------------
MAX_ORGANIZE_PAGES = 100
MAX_ORGANIZE_SOURCES = 10


def pdf_thumbnails(src: Path, width: int = 220) -> list[dict]:
    try:
        doc = pymupdf.open(src)
    except Exception:
        raise ValueError("This file is not a valid PDF")
    try:
        if doc.needs_pass:
            raise ValueError("This PDF is password protected. Unlock it first")
        if doc.page_count > MAX_ORGANIZE_PAGES:
            raise ValueError(
                f"Too many pages (max {MAX_ORGANIZE_PAGES}). Split the PDF first"
            )
        thumbs = []
        for page in doc:
            w, h = max(page.rect.width, 1), max(page.rect.height, 1)
            zoom = min(width / w, width * 1.5 / h)
            pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
            thumbs.append({
                "img": base64.b64encode(pix.tobytes("jpg", jpg_quality=60)).decode(),
                "w": round(w, 1),
                "h": round(h, 1),
            })
        return thumbs
    finally:
        doc.close()


def organize_pdf(streams: list, plan_json: str, out: Path) -> None:
    try:
        plan = json.loads(plan_json)
    except ValueError:
        raise ValueError("Invalid page plan")
    if not isinstance(plan, list) or not plan:
        raise ValueError("Add at least one page")
    if len(plan) > MAX_ORGANIZE_PAGES * 2:
        raise ValueError("Too many pages in the result")
    if not 1 <= len(streams) <= MAX_ORGANIZE_SOURCES:
        raise ValueError(f"Use between 1 and {MAX_ORGANIZE_SOURCES} files")

    readers = []
    for s in streams:
        try:
            r = PdfReader(s)
        except Exception:
            raise ValueError("One of the files is not a valid PDF")
        if r.is_encrypted:
            raise ValueError("A PDF is password protected. Unlock it first")
        if len(r.pages) == 0:
            raise ValueError("A PDF has no pages")
        readers.append(r)

    first = readers[0].pages[0]
    size = (float(first.mediabox.width), float(first.mediabox.height))

    writer = PdfWriter()
    for item in plan:
        if not isinstance(item, dict):
            raise ValueError("Invalid page plan")
        rot = item.get("rotate", 0)
        if isinstance(rot, bool) or rot not in (0, 90, 180, 270):
            raise ValueError("Rotation must be 0, 90, 180 or 270")

        if item.get("blank"):
            new = writer.add_blank_page(*size)  # same size as the previous real page
        else:
            src, page = item.get("src"), item.get("page")
            if isinstance(src, bool) or not isinstance(src, int) or not 0 <= src < len(readers):
                raise ValueError("Invalid page plan")
            total = len(readers[src].pages)
            if isinstance(page, bool) or not isinstance(page, int) or not 1 <= page <= total:
                raise ValueError(f"Pages must be between 1 and {total}")
            source_page = readers[src].pages[page - 1]
            size = (float(source_page.mediabox.width), float(source_page.mediabox.height))
            new = writer.add_page(source_page)
        if rot:
            new.rotate(rot)  # rotate the copy, so duplicates stay independent

    with out.open("wb") as fh:
        writer.write(fh)


# ---------------------------------------------------------------
# Shared helpers for Sign, Compare and Redact
# ---------------------------------------------------------------
MAX_SIGN_PAGES = 30
MAX_SIGNATURES = 50
MAX_SIGN_IMAGES = 10
MAX_COMPARE_PAGES = 30


def _open_checked(path: Path):
    try:
        doc = pymupdf.open(path)
    except Exception:
        raise ValueError("This file is not a valid PDF")
    if doc.needs_pass:
        doc.close()
        raise ValueError("A PDF is password protected. Unlock it first")
    return doc


def pdf_pages(src: Path, width: int = 700) -> list[dict]:
    doc = _open_checked(src)
    try:
        if doc.page_count > MAX_SIGN_PAGES:
            raise ValueError(f"Too many pages (max {MAX_SIGN_PAGES}). Extract fewer pages first")
        pages = []
        for page in doc:
            zoom = width / max(page.rect.width, 1)
            pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
            pages.append({
                "img": base64.b64encode(pix.tobytes("jpg", jpg_quality=75)).decode(),
                "w": round(page.rect.width, 1),
                "h": round(page.rect.height, 1),
            })
        return pages
    finally:
        doc.close()


def _num(v, lo, hi):
    if isinstance(v, bool) or not isinstance(v, (int, float)) or not lo <= v <= hi:
        raise ValueError("Invalid position")
    return float(v)


def _render(page, width):
    zoom = width / max(page.rect.width, 1)
    pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
    return Image.frombytes("RGB", (pix.width, pix.height), pix.samples), pymupdf.Matrix(zoom, zoom)


def _words(page):
    return [(pymupdf.Rect(w[:4]) * page.rotation_matrix, w[4]) for w in page.get_text("words")]


def _highlight(img, rects, color):
    layer = Image.new("RGBA", img.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(layer)
    for r in rects:
        draw.rectangle([r.x0, r.y0, r.x1, r.y1], fill=color)
    return Image.alpha_composite(img.convert("RGBA"), layer).convert("RGB")


def _b64(img) -> str:
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=70)
    return base64.b64encode(buf.getvalue()).decode()


# ---------------------------------------------------------------
# Sign PDF
# ---------------------------------------------------------------
def _prepare_signature(data: bytes):
    try:
        img = Image.open(io.BytesIO(data))
        if img.width * img.height > 4_000_000:
            raise ValueError("Signature image is too large")
        img.load()
    except ValueError:
        raise
    except Exception:
        raise ValueError("The signature image is not valid")
    img = img.convert("RGBA")
    box = img.getchannel("A").getbbox()
    if not box:
        raise ValueError("A signature is empty")
    return img.crop(box)


def sign_pdf(src: Path, images: list[bytes], placements_json: str, out: Path) -> None:
    try:
        items = json.loads(placements_json)
    except ValueError:
        raise ValueError("Invalid placement data")
    if not isinstance(items, list) or not items:
        raise ValueError("Place something on a page first")
    if len(items) > MAX_SIGNATURES:
        raise ValueError(f"Too many items (max {MAX_SIGNATURES})")
    if not 1 <= len(images) <= MAX_SIGN_IMAGES:
        raise ValueError(f"Use between 1 and {MAX_SIGN_IMAGES} different items")

    sigs = [_prepare_signature(b) for b in images]

    doc = _open_checked(src)
    try:
        cache = {}
        for it in items:
            if not isinstance(it, dict):
                raise ValueError("Invalid placement data")
            pno, idx = it.get("page"), it.get("img")
            if isinstance(pno, bool) or not isinstance(pno, int) or not 1 <= pno <= doc.page_count:
                raise ValueError(f"Pages must be between 1 and {doc.page_count}")
            if isinstance(idx, bool) or not isinstance(idx, int) or not 0 <= idx < len(sigs):
                raise ValueError("Invalid placement data")
            xf, yf = _num(it.get("x"), 0, 1), _num(it.get("y"), 0, 1)
            wf = _num(it.get("w"), 0.02, 1)

            sig = sigs[idx]
            page = doc[pno - 1]
            W, H = page.rect.width, page.rect.height
            w = wf * W
            h = w * sig.height / sig.width
            if h > H:
                h, w = H, H * sig.width / sig.height
            x0 = min(max(xf * W, 0), W - w)
            y0 = min(max(yf * H, 0), H - h)
            rect = pymupdf.Rect(x0, y0, x0 + w, y0 + h) * page.derotation_matrix
            rect.normalize()

            rot = page.rotation
            key = (idx, rot)
            if key not in cache:  # turn the image so it looks upright on rotated pages
                img = sig.rotate(rot, expand=True) if rot else sig
                buf = io.BytesIO()
                img.save(buf, "PNG")
                cache[key] = buf.getvalue()
            page.insert_image(rect, stream=cache[key])
        doc.save(out, garbage=3, deflate=True)
    finally:
        doc.close()


# ---------------------------------------------------------------
# Compare PDF
# ---------------------------------------------------------------
def compare_pdfs(path_a: Path, path_b: Path, width: int = 500) -> dict:
    doc_a, doc_b = _open_checked(path_a), None
    try:
        doc_b = _open_checked(path_b)
        na, nb = doc_a.page_count, doc_b.page_count
        if max(na, nb) > MAX_COMPARE_PAGES:
            raise ValueError(f"Too many pages (max {MAX_COMPARE_PAGES} per file)")

        pages, removed_total, added_total, changed_pages = [], 0, 0, 0
        for i in range(max(na, nb)):
            entry = {"page": i + 1, "a": None, "b": None,
                     "removed": 0, "added": 0, "visual": 0.0}
            pa = doc_a[i] if i < na else None
            pb = doc_b[i] if i < nb else None

            if pa and pb:
                img_a, mat_a = _render(pa, width)
                img_b, mat_b = _render(pb, width)
                wa, wb = _words(pa), _words(pb)
                sm = difflib.SequenceMatcher(
                    None, [t for _, t in wa], [t for _, t in wb], autojunk=False
                )
                rem, add = [], []
                for tag, i1, i2, j1, j2 in sm.get_opcodes():
                    if tag in ("replace", "delete"):
                        rem += [wa[k][0] * mat_a for k in range(i1, i2)]
                    if tag in ("replace", "insert"):
                        add += [wb[k][0] * mat_b for k in range(j1, j2)]
                diff = ImageChops.difference(
                    img_a.convert("L"), img_b.resize(img_a.size).convert("L")
                ).point(lambda p: 255 if p > 40 else 0)
                visual = 100 * diff.histogram()[255] / (diff.width * diff.height)
                entry.update(
                    removed=len(rem), added=len(add), visual=round(visual, 2),
                    a=_b64(_highlight(img_a, rem, (255, 0, 0, 90))),
                    b=_b64(_highlight(img_b, add, (0, 170, 0, 90))),
                )
                entry["changed"] = bool(rem or add or visual > 0.5)
            else:
                page, key = (pa, "a") if pa else (pb, "b")
                entry[key] = _b64(_render(page, width)[0])
                entry["changed"] = True
                entry["missing"] = "b" if pa else "a"

            removed_total += entry["removed"]
            added_total += entry["added"]
            changed_pages += entry["changed"]
            pages.append(entry)

        return {"pages_a": na, "pages_b": nb, "changed_pages": changed_pages,
                "words_removed": removed_total, "words_added": added_total, "pages": pages}
    finally:
        doc_a.close()
        if doc_b:
            doc_b.close()


# ---------------------------------------------------------------
# Redact PDF
# ---------------------------------------------------------------
MAX_REDACT_PAGES = 30
MAX_REDACT_TERMS = 50
MAX_REDACT_BOXES = 500
MAX_REDACT_AREAS = 5000

REDACT_PATTERNS = {
    "email": r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}",
    "phone": r"(?<!\d)\+?\d[\d\s().\-]{7,}\d(?!\d)",
    "url": r"(?:https?://|www\.)[^\s]+",
    "card": r"(?<!\d)(?:\d[ \-]?){13,16}(?!\d)",
}


def redact_pdf(src: Path, terms_text: str, pattern_keys: str, boxes_json: str, out: Path) -> int:
    terms = [t.strip() for t in terms_text.splitlines() if t.strip()]
    if len(terms) > MAX_REDACT_TERMS:
        raise ValueError(f"Too many search terms (max {MAX_REDACT_TERMS})")
    if any(len(t) > 200 for t in terms):
        raise ValueError("A search term is too long (max 200 characters)")

    keys = [k.strip() for k in pattern_keys.split(",") if k.strip()]
    if any(k not in REDACT_PATTERNS for k in keys):
        raise ValueError("Unknown pattern")

    try:
        boxes = json.loads(boxes_json or "[]")
    except ValueError:
        raise ValueError("Invalid box data")
    if not isinstance(boxes, list) or len(boxes) > MAX_REDACT_BOXES:
        raise ValueError(f"Too many boxes (max {MAX_REDACT_BOXES})")

    if not terms and not keys and not boxes:
        raise ValueError("Draw a box, add a search term or choose a pattern first")

    doc = _open_checked(src)
    try:
        if doc.page_count > MAX_REDACT_PAGES:
            raise ValueError(f"Too many pages (max {MAX_REDACT_PAGES}). Extract fewer pages first")

        by_page = {}
        for b in boxes:
            if not isinstance(b, dict):
                raise ValueError("Invalid box data")
            pno = b.get("page")
            if isinstance(pno, bool) or not isinstance(pno, int) or not 1 <= pno <= doc.page_count:
                raise ValueError(f"Pages must be between 1 and {doc.page_count}")
            by_page.setdefault(pno, []).append(b)

        total = 0
        for pno, page in enumerate(doc, start=1):
            rects = []
            for t in terms:
                rects += page.search_for(t)

            if keys:
                text = page.get_text("text")
                found = set()
                for k in keys:
                    found.update(m.group(0).strip() for m in re.finditer(REDACT_PATTERNS[k], text))
                for s in found:
                    rects += page.search_for(s)

            W, H = page.rect.width, page.rect.height  # size as displayed (rotation included)
            for b in by_page.get(pno, []):
                x, y = _num(b.get("x"), 0, 1), _num(b.get("y"), 0, 1)
                w, h = _num(b.get("w"), 0.001, 1), _num(b.get("h"), 0.001, 1)
                r = pymupdf.Rect(x * W, y * H, min(x + w, 1) * W, min(y + h, 1) * H)
                r = r * page.derotation_matrix  # displayed -> page coordinates
                r.normalize()
                rects.append(r)

            if not rects:
                continue
            total += len(rects)
            if total > MAX_REDACT_AREAS:
                raise ValueError("Too many areas to redact. Use more specific terms")
            for r in rects:
                page.add_redact_annot(r, fill=(0, 0, 0))
            page.apply_redactions(images=2)  # 2 = also blank image pixels under the box

        if total == 0:
            raise ValueError("Nothing found to redact. Check your search terms or draw a box")

        doc.set_metadata({})
        doc.del_xml_metadata()
        doc.save(out, garbage=4, deflate=True, clean=True)
        return total
    finally:
        doc.close()


# ---------------------------------------------------------------
# Crop PDF
# ---------------------------------------------------------------
def crop_pdf(stream, box_json: str, spec: str, out: Path) -> None:
    try:
        box = json.loads(box_json)
    except ValueError:
        raise ValueError("Invalid crop area")
    if not isinstance(box, dict):
        raise ValueError("Invalid crop area")
    u0, v0 = _num(box.get("x"), 0, 1), _num(box.get("y"), 0, 1)
    u1 = min(u0 + _num(box.get("w"), 0.01, 1), 1)
    v1 = min(v0 + _num(box.get("h"), 0.01, 1), 1)

    reader = PdfReader(stream)
    if reader.is_encrypted:
        raise ValueError("This PDF is password protected. Unlock it first")
    total = len(reader.pages)
    targets = set(_safe_pages(spec, total)) if spec.strip() else set(range(total))

    writer = PdfWriter()
    for i, src_page in enumerate(reader.pages):
        page = writer.add_page(src_page)
        if i not in targets:
            continue
        cb = page.cropbox  # the visible area (falls back to the media box)
        left, bottom, right, top = float(cb.left), float(cb.bottom), float(cb.right), float(cb.top)
        bw, bh = right - left, top - bottom
        rot = page.rotation % 360

        def to_pdf(u, v):  # displayed fractions -> PDF coordinates (y up)
            if rot == 90:
                return left + v * bw, bottom + u * bh
            if rot == 180:
                return left + (1 - u) * bw, bottom + v * bh
            if rot == 270:
                return left + (1 - v) * bw, top - u * bh
            return left + u * bw, top - v * bh

        (xa, ya), (xb, yb) = to_pdf(u0, v0), to_pdf(u1, v1)
        page.cropbox = RectangleObject([min(xa, xb), min(ya, yb), max(xa, xb), max(ya, yb)])

    with out.open("wb") as fh:
        writer.write(fh)


# ---------------------------------------------------------------
# Edit PDF
# ---------------------------------------------------------------
MAX_EDIT_PAGES = 20
MAX_EDIT_OPS = 300
MAX_EDIT_LINES = 3000
EDIT_TYPES = {"text", "replace", "whiteout", "highlight", "rect", "ellipse", "line", "draw"}
_FONT_SETS = (
    ("helv", "hebo", "heit", "hebi"),  # sans
    ("tiro", "tibo", "tiit", "tibi"),  # serif
    ("cour", "cobo", "coit", "cobi"),  # mono
)
EDIT_FONTS = {f for fs in _FONT_SETS for f in fs}
_HEX = re.compile(r"^#[0-9a-fA-F]{6}$")


def _rgb01(c):
    if not isinstance(c, str) or not _HEX.match(c):
        raise ValueError("Invalid colour")
    return tuple(int(c[i:i + 2], 16) / 255 for i in (1, 3, 5))


def _font_for(name: str, flags: int) -> str:
    low = name.lower()
    bold = bool(flags & 16) or "bold" in low
    italic = bool(flags & 2) or "italic" in low or "oblique" in low
    if (flags & 8) or "mono" in low or "courier" in low:
        fam = _FONT_SETS[2]
    elif (flags & 4) or "times" in low or ("serif" in low and "sans" not in low):
        fam = _FONT_SETS[1]
    else:
        fam = _FONT_SETS[0]
    return fam[(1 if bold else 0) + (2 if italic else 0)]


def pdf_edit_info(src: Path, width: int = 700) -> list[dict]:
    doc = _open_checked(src)
    try:
        if doc.page_count > MAX_EDIT_PAGES:
            raise ValueError(f"Too many pages (max {MAX_EDIT_PAGES}). Extract fewer pages first")
        pages, total = [], 0
        for page in doc:
            W, H = max(page.rect.width, 1), max(page.rect.height, 1)
            zoom = width / W
            pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
            lines = []
            for b in page.get_text("dict")["blocks"]:
                if b["type"] != 0:
                    continue
                for ln in b["lines"]:
                    spans = [s for s in ln["spans"] if s["text"].strip()]
                    if not spans or total >= MAX_EDIT_LINES:
                        continue
                    if page.rotation == 0 and abs(ln["dir"][1]) > 0.1:
                        continue  # skip text that is not horizontal
                    s0 = spans[0]
                    r = pymupdf.Rect(ln["bbox"]) * page.rotation_matrix
                    r.normalize()
                    o = pymupdf.Point(s0["origin"]) * page.rotation_matrix
                    lines.append({
                        "x": round(r.x0 / W, 5), "y": round(r.y0 / H, 5),
                        "w": round(r.width / W, 5), "h": round(r.height / H, 5),
                        "ox": round(o.x / W, 5), "oy": round(o.y / H, 5),
                        "text": "".join(s["text"] for s in ln["spans"]).strip(),
                        "size": round(s0["size"], 1),
                        "color": "#%06x" % s0["color"],
                        "font": _font_for(s0["font"], s0["flags"]),
                    })
                    total += 1
            pages.append({
                "img": base64.b64encode(pix.tobytes("jpg", jpg_quality=75)).decode(),
                "w": round(W, 1), "h": round(H, 1), "lines": lines,
            })
        return pages
    finally:
        doc.close()


def _op_rect(page, op, shrink=0.0):
    W, H = page.rect.width, page.rect.height
    x, y = _num(op.get("x"), 0, 1), _num(op.get("y"), 0, 1)
    w, h = _num(op.get("w"), 0.001, 1), _num(op.get("h"), 0.001, 1)
    y0, y1 = (y + h * shrink) * H, (min(y + h, 1) - h * shrink) * H
    r = pymupdf.Rect(x * W, y0, min(x + w, 1) * W, y1) * page.derotation_matrix
    r.normalize()
    return r


def _op_point(page, fx, fy):
    return pymupdf.Point(
        _num(fx, 0, 1) * page.rect.width, _num(fy, 0, 1) * page.rect.height
    ) * page.derotation_matrix


def edit_pdf(src: Path, ops_json: str, out: Path) -> None:
    try:
        ops = json.loads(ops_json)
    except ValueError:
        raise ValueError("Invalid edit data")
    if not isinstance(ops, list) or not ops:
        raise ValueError("Make a change first")
    if len(ops) > MAX_EDIT_OPS:
        raise ValueError(f"Too many changes (max {MAX_EDIT_OPS})")

    latin = pymupdf.Font("helv")

    def check_text(t):
        if not isinstance(t, str) or len(t) > 5000:
            raise ValueError("Invalid text")
        for ch in t:
            if ch not in "\n\r\t " and not latin.has_glyph(ord(ch)):
                raise ValueError(
                    f"The character '{ch}' is not supported yet. "
                    "Use Latin letters, numbers and common symbols"
                )

    doc = _open_checked(src)
    try:
        if doc.page_count > MAX_EDIT_PAGES:
            raise ValueError(f"Too many pages (max {MAX_EDIT_PAGES})")
        by_page = {}
        for op in ops:
            if not isinstance(op, dict) or op.get("type") not in EDIT_TYPES:
                raise ValueError("Invalid edit data")
            pno = op.get("page")
            if isinstance(pno, bool) or not isinstance(pno, int) or not 1 <= pno <= doc.page_count:
                raise ValueError(f"Pages must be between 1 and {doc.page_count}")
            by_page.setdefault(pno, []).append(op)

        for pno, plist in by_page.items():
            page = doc[pno - 1]
            rot = page.rotation

            # 1. Remove what is under whiteouts and replaced lines
            covers = [o for o in plist if o["type"] in ("whiteout", "replace")]
            for o in covers:
                page.add_redact_annot(
                    _op_rect(page, o, shrink=0.1 if o["type"] == "replace" else 0.0),
                    fill=(1, 1, 1),
                )
            if covers:
                page.apply_redactions(images=2)

            # 2. Shapes
            shape, drew = page.new_shape(), False
            for o in plist:
                t = o["type"]
                if t == "highlight":
                    shape.draw_rect(_op_rect(page, o))
                    shape.finish(color=None, fill=_rgb01(o.get("color", "#ffe600")), fill_opacity=0.35)
                    drew = True
                elif t in ("rect", "ellipse"):
                    col, sw = _rgb01(o.get("color", "#000000")), _num(o.get("sw", 1), 0.25, 20)
                    r = _op_rect(page, o)
                    shape.draw_rect(r) if t == "rect" else shape.draw_oval(r)
                    shape.finish(color=col, fill=col if o.get("fill") is True else None, width=sw)
                    drew = True
                elif t == "line":
                    col, sw = _rgb01(o.get("color", "#000000")), _num(o.get("sw", 1), 0.25, 20)
                    shape.draw_line(_op_point(page, o.get("x1"), o.get("y1")),
                                    _op_point(page, o.get("x2"), o.get("y2")))
                    shape.finish(color=col, width=sw, closePath=False)
                    drew = True
                elif t == "draw":
                    col, sw = _rgb01(o.get("color", "#000000")), _num(o.get("sw", 1), 0.25, 20)
                    pts = o.get("pts")
                    if not isinstance(pts, list) or not 2 <= len(pts) <= 2000:
                        raise ValueError("Invalid drawing")
                    points = []
                    for p in pts:
                        if not isinstance(p, list) or len(p) != 2:
                            raise ValueError("Invalid drawing")
                        points.append(_op_point(page, p[0], p[1]))
                    shape.draw_polyline(points)
                    shape.finish(color=col, width=sw, closePath=False, lineCap=1, lineJoin=1)
                    drew = True
            if drew:
                shape.commit()

            # 3. Text on top
            for o in plist:
                if o["type"] not in ("text", "replace"):
                    continue
                txt = o.get("text", "")
                check_text(txt)
                if not txt.strip():
                    continue  # an empty replacement just removes the line
                font = o.get("font", "helv")
                if font not in EDIT_FONTS:
                    raise ValueError("Unknown font")
                size = _num(o.get("size", 12), 4, 200)
                col = _rgb01(o.get("color", "#000000"))
                if o["type"] == "text":
                    rc = page.insert_textbox(_op_rect(page, o), txt, fontsize=size,
                                             fontname=font, color=col, rotate=rot)
                    if rc < 0:
                        raise ValueError(
                            f"Text does not fit its box on page {pno}. "
                            "Make the box bigger or the text smaller"
                        )
                else:
                    page.insert_text(_op_point(page, o.get("ox"), o.get("oy")),
                                     txt.replace("\n", " "), fontsize=size,
                                     fontname=font, color=col, rotate=rot)

        doc.save(out, garbage=3, deflate=True)
    finally:
        doc.close()


# ---------------------------------------------------------------
# Page numbers and watermark
# ---------------------------------------------------------------
MAX_STAMP_PAGES = 300
NUM_FORMATS = {
    "n": "{n}",
    "page_n": "Page {n}",
    "n_of_t": "{n} / {t}",
    "page_n_of_t": "Page {n} of {t}",
}
NUM_POSITIONS = {"bl", "bc", "br", "tl", "tc", "tr"}
ANGLE_SIGN = 1  # change to -1 if the diagonal watermark leans the wrong way


def _check_latin(text: str) -> None:
    latin = pymupdf.Font("helv")
    for ch in text:
        if not latin.has_glyph(ord(ch)):
            raise ValueError(
                f"The character '{ch}' is not supported yet. "
                "Use Latin letters, numbers and common symbols"
            )


def _int_in(value, lo, hi, name):
    try:
        v = int(value)
    except (TypeError, ValueError):
        raise ValueError(f"{name} must be a whole number")
    if not lo <= v <= hi:
        raise ValueError(f"{name} must be between {lo} and {hi}")
    return v


def _target_pages(doc, spec: str) -> list[int]:
    total = doc.page_count
    if total > MAX_STAMP_PAGES:
        raise ValueError(f"Too many pages (max {MAX_STAMP_PAGES}). Split the PDF first")
    if spec.strip():
        return sorted(set(_safe_pages(spec, total)))
    return list(range(total))


def add_page_numbers(src: Path, position: str, fmt: str, start: str, size: str,
                     margin: str, color: str, spec: str, out: Path) -> None:
    if position not in NUM_POSITIONS:
        raise ValueError("Unknown position")
    if fmt not in NUM_FORMATS:
        raise ValueError("Unknown number format")
    first = _int_in(start, 0, 100000, "Start number")
    fs = _int_in(size, 6, 36, "Size")
    mg = _int_in(margin, 0, 150, "Margin")
    col = _rgb01(color)

    doc = _open_checked(src)
    try:
        targets = _target_pages(doc, spec)
        last = first + len(targets) - 1  # the "total" in "1 of N"
        for k, i in enumerate(targets):
            page = doc[i]
            text = NUM_FORMATS[fmt].format(n=first + k, t=last)
            W, H = page.rect.width, page.rect.height  # size as displayed
            tw = pymupdf.get_text_length(text, fontname="helv", fontsize=fs)
            if position[1] == "l":
                x = mg
            elif position[1] == "c":
                x = (W - tw) / 2
            else:
                x = W - mg - tw
            y = mg + fs if position[0] == "t" else H - mg
            pt = pymupdf.Point(x, y) * page.derotation_matrix
            page.insert_text(pt, text, fontsize=fs, fontname="helv",
                             color=col, rotate=page.rotation)
        doc.save(out, garbage=3, deflate=True)
    finally:
        doc.close()


def add_watermark(src: Path, text: str, size: str, color: str, opacity: str,
                  angle: str, layout: str, position: str, spec: str, out: Path) -> None:
    text = " ".join(text.split())
    if not text:
        raise ValueError("Enter the watermark text")
    if len(text) > 60:
        raise ValueError("The watermark text is too long (max 60 characters)")
    _check_latin(text)
    fs0 = _int_in(size, 10, 200, "Size")
    pct = _int_in(opacity, 5, 100, "Opacity")
    deg = _int_in(angle, -90, 90, "Angle")
    if layout not in ("center", "tile"):
        raise ValueError("Unknown layout")
    if position not in ("over", "behind"):
        raise ValueError("Unknown position")
    col = _rgb01(color)

    doc = _open_checked(src)
    try:
        targets = _target_pages(doc, spec)
        for i in targets:
            page = doc[i]
            W, H = page.rect.width, page.rect.height
            fs = fs0
            tw = pymupdf.get_text_length(text, fontname="hebo", fontsize=fs)

            if layout == "center":
                maxw = 0.85 * W if deg == 0 else 0.7 * math.hypot(W, H)
                if tw > maxw:  # shrink so the text stays on the page
                    fs, tw = fs * maxw / tw, maxw
                centers = [(W / 2, H / 2)]
            else:
                stepx = max(tw * (1.1 if deg else 0.4) + fs * 2, W / 6)
                stepy = max(tw * 0.7 if deg else fs * 5, H / 10)
                centers, row, cy = [], 0, stepy / 2
                while cy < H + stepy:
                    cx = stepx / 2 + (stepx / 2 if row % 2 else 0) - stepx
                    while cx < W + stepx:
                        centers.append((cx, cy))
                        cx += stepx
                    cy += stepy
                    row += 1

            for cx, cy in centers:
                start = pymupdf.Point(cx - tw / 2, cy + fs * 0.35) * page.derotation_matrix
                pivot = pymupdf.Point(cx, cy) * page.derotation_matrix
                kw = dict(fontsize=fs, fontname="hebo", color=col,
                          fill_opacity=pct / 100, stroke_opacity=pct / 100,
                          rotate=page.rotation, overlay=(position == "over"))
                if deg:
                    kw["morph"] = (pivot, pymupdf.Matrix(ANGLE_SIGN * deg))
                page.insert_text(start, text, **kw)
        doc.save(out, garbage=3, deflate=True)
    finally:
        doc.close()
